# closing pass (c): add the section 20v sdr-elist fallback to every sensor-get row that lacks it.
# Std format (REVIEW_WORKFLOW_LOGIC.md section 20v):
#   ipmitool ... sensor get 'X' 2>&1; echo ---SDR-FALLBACK---; ipmitool ... sdr elist 2>&1 | head -100
#
# Three shapes, located by exact regex with ABORT-on-any-unexpected-row:
#   A) OOB-lanplus direct (agent host): ipmitool -I lanplus -C 17 -H "$BMC_IP" -U "$BMC_USER" -P "$BMC_PASS" sensor get '<N>' 2>&1
#   B) ssh-noipmi bare (10): ssh ... "$DUT_IP" "sensor get '<N>'"
#        -> ALSO fixes R29 (bare `sensor get` without ipmitool) by adding the ipmitool prefix in-band.
#   C) ssh-inband-ipmitool (1): ssh ... "$DUT_IP" "sudo ipmitool sensor get '<N>' 2>&1"
import openpyxl, re, sys

XLSX = 'data/REVISED_commands_merged_with_raw.xlsx'
SHEETS = ['Functionality', 'Reliability', 'Performance', 'Compatibility', 'Stability', '(No Main Function)']

PREFIX_OOB = 'ipmitool -I lanplus -C 17 -H "$BMC_IP" -U "$BMC_USER" -P "$BMC_PASS"'
RE_A = re.compile(r'(%s sensor get \'[^\']*\' 2>&1)' % re.escape(PREFIX_OOB))
RE_A2 = re.compile(r'(%s sensor get \'[^\']*\')$' % re.escape(PREFIX_OOB))  # e.g. sensor get '*' (no 2>&1)
RE_B = re.compile(r'(sshpass -p "\$DUT_PASS" ssh -o StrictHostKeyChecking=no \$DUT_USER@\$DUT_IP "sensor get \'[^\']*\')(\")?')
RE_C = re.compile(r'(sshpass -p "\$DUT_PASS" ssh -o StrictHostKeyChecking=no \$DUT_USER@\$DUT_IP "sudo ipmitool sensor get \'[^\']*\' 2>&1)(\")?')

wb = openpyxl.load_workbook(XLSX)
changed = {}
for sn in SHEETS:
    ws = wb[sn]
    hdr = [c.value for c in ws[1]]
    ci = {h: i for i, h in enumerate(hdr)}
    for r in range(2, ws.max_row + 1):
        aic = ws.cell(row=r, column=ci['ai_commands'] + 1).value
        if not isinstance(aic, str):
            continue
        if 'sdr elist' in aic or 'sensor list' in aic:
            continue
        new = None
        if RE_A.search(aic):
            # direct OOB: append same-prefix sdr elist after the sensor-get 2>&1
            def repA(m):
                cmd = m.group(1)
                return cmd + '; echo ---SDR-FALLBACK---; %s sdr elist 2>&1 | head -100' % PREFIX_OOB
            new, n = RE_A.subn(repA, aic)
            if n != 1:
                sys.exit('ABORT %s r%d: A matched %d times' % (sn, r, n))
        elif RE_A2.search(aic):
            # direct OOB bulk (sensor get '*' without 2>&1): append 2>&1 + fallback
            def repA2(m):
                cmd = m.group(1)
                return cmd + ' 2>&1; echo ---SDR-FALLBACK---; %s sdr elist 2>&1 | head -100' % PREFIX_OOB
            new, n = RE_A2.subn(repA2, aic)
            if n != 1:
                sys.exit('ABORT %s r%d: A2 matched %d times' % (sn, r, n))
        elif RE_B.search(aic):
            # bare in-band sensor get: add ipmitool prefix + fallback inside the ssh remote
            def repB(m):
                name = re.search(r"sensor get '([^']*)'", m.group(1)).group(1)
                return ('sshpass -p "$DUT_PASS" ssh -o StrictHostKeyChecking=no $DUT_USER@$DUT_IP '
                        '"ipmitool sensor get \'%s\' 2>&1; echo ---SDR-FALLBACK---; '
                        'ipmitool sdr elist 2>&1 | head -100"' % name)
            new, n = RE_B.subn(repB, aic)
            if n != 1:
                sys.exit('ABORT %s r%d: B matched %d times' % (sn, r, n))
        elif RE_C.search(aic):
            def repC(m):
                return m.group(1) + '; echo ---SDR-FALLBACK---; sudo ipmitool sdr elist 2>&1 | head -100"'
            new, n = RE_C.subn(repC, aic)
            if n != 1:
                sys.exit('ABORT %s r%d: C matched %d times' % (sn, r, n))
        if new is not None and new != aic:
            ws.cell(row=r, column=ci['ai_commands'] + 1).value = new
            changed[(sn, r)] = (aic, new)

print('rows with sdr-elist fallback added:', len(changed))
from collections import Counter
print('by sheet:', dict(Counter(k[0] for k in changed)))

DRY = len(sys.argv) > 1 and sys.argv[1] == 'dry'
if DRY:
    print('DRY - not saved')
else:
    wb.save(XLSX)
    print('saved', XLSX)
