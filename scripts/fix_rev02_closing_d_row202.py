# closing pass (d) + row-202:
#  d) r1957/r2090 `bad_page_threshold=<val>` -> `${BAD_PAGE_THRESHOLD:?...}` (R7 vendor slot)
#  row202 Wistron-HW-00211 Hot Plugging: Q-LIT - unescaped inner double-quotes inside the ssh
#      remote (`watch -n1 "lsusb 2>&1"`, `grep -iE "usb|..."`) -> single-quote the inner literals.
import openpyxl, sys

XLSX = 'data/REVISED_commands_merged_with_raw.xlsx'
SHEET = 'Functionality'

FIX_D = {
    1957: (
        'bad_page_threshold=<val>',
        'bad_page_threshold=${BAD_PAGE_THRESHOLD:?operator provides the amdgpu RAS bad-page threshold value}',
    ),
    2090: (
        'bad_page_threshold=<val>',
        'bad_page_threshold=${BAD_PAGE_THRESHOLD:?operator provides the amdgpu RAS bad-page threshold value}',
    ),
}

NEW_ROW202 = ('-- Hot Plugging: verify USB hot-plugging (insert/remove a USB key without shutdown). '
              "Operator physically plugs/unplugs the key into the host USB ports; agent watches the OS USB event + device state "
              'sshpass -p "$DUT_PASS" ssh -o StrictHostKeyChecking=no $DUT_USER@$DUT_IP '
              '"watch -n1 \'lsusb 2>&1\' 2>&1 & sleep 6; kill %1; '
              "dmesg -T 2>&1 | grep -iE 'usb|new.*device|USB disconnect' | tail -20 2>&1\" "
              'to confirm detect/remove. -- USB hot-plugging instance; verify detect/remove without shutdown on each rear USB port.')

wb = openpyxl.load_workbook(XLSX)
ws = wb[SHEET]

for row, (old, new) in FIX_D.items():
    cur = ws.cell(row=row, column=15).value or ''
    if old not in cur:
        sys.exit('ABORT r%d: %r not present in col15' % (row, old))
    ws.cell(row=row, column=15).value = cur.replace(old, new, 1)
    print('r%d: <val> -> ${BAD_PAGE_THRESHOLD:?...}' % row)

# row 202 - whole-cell rewrite with Q-LIT fix
cur202 = ws.cell(row=202, column=15).value or ''
markers = ['Hot Plugging', 'watch -n1', '$DUT_IP "']
if any(m not in cur202 for m in markers):
    sys.exit('ABORT r202: expected markers not present; current=%r' % cur202)
ws.cell(row=202, column=15).value = NEW_ROW202
print('r202: Q-LIT fixed (inner quotes -> single)')

print('saved', XLSX)
wb.save(XLSX)
