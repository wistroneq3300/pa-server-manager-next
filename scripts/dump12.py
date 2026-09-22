# Dump rows 2203-2292 (item 2201-2290) to /tmp/rows2203_2292.json for the batch12 report generator.
import openpyxl, json
p="data/REVISED_commands_merged_with_raw.xlsx"
wb=openpyxl.load_workbook(p, read_only=True, data_only=True)
print("sheets:", wb.sheetnames)
ws=wb["Functionality"]
print("Functionality max_row:", ws.max_row)
print("(reads the current post-batch12-fix xlsx -> /tmp/rows2203_2292.json)")
HDR=["Code","Sub Function","Test Set","Items","Procedure","Criteria",
     "Bundle","Branch","Script","Attended Time","Machine Time","Latest Updated",
     "ai_can_execute","ai_packages_needed","ai_commands","ai_logs_output","risk","remark"]
rows=[]
for i,row in enumerate(ws.iter_rows(values_only=True, min_row=2), start=2):
    if not (2203 <= i <= 2292): continue
    rec={"_row": i}
    for k,name in enumerate(HDR):
        rec[name] = row[k] if k < len(row) else None
    rows.append(rec)
wb.close()
json.dump(rows, open('/tmp/rows2203_2292.json','w'), ensure_ascii=False, indent=1)
print("wrote", len(rows), "rows; rows 2203-2292")
