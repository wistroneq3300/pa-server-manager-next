# Precise per-(sheet,row) ai_commands fix for Round-02 batch 5 (Functionality rows 803-1002, item 801-1000).
# R29  : bare ipmitool subcommand missing the `ipmitool` keyword inside DUT-ssh -> add `sudo ipmitool` (+ raw-byte fix where wrong).
# R5   : OOB ipmitool lanplus / BMC redfish curl wrongly wrapped in DUT-ssh with broken nested quotes -> move to agent-host.
# Q-LIT: inner grep double-quoted pattern -> single quotes inside ssh.
# Pure ASCII only in this script. Uses python-openpyxl.
import openpyxl, sys

XLSX = 'data/REVISED_commands_merged_with_raw.xlsx'
SHEET = 'Functionality'

FIX = {
# --- R29: bare ipmitool subcommand (missing keyword) on DUT ---
986: r'''sshpass -p "$DUT_PASS" ssh -o StrictHostKeyChecking=no $DUT_USER@$DUT_IP "sudo ipmitool chassis power cycle 2>&1; sleep 10; sudo ipmitool chassis power status 2>&1"''',
992: r'''sshpass -p "$DUT_PASS" ssh -o StrictHostKeyChecking=no $DUT_USER@$DUT_IP "sudo ipmitool raw 0x06 0x01 2>&1"''',
993: r'''sshpass -p "$DUT_PASS" ssh -o StrictHostKeyChecking=no $DUT_USER@$DUT_IP "sudo ipmitool raw 0x06 0x02 2>&1; sleep 90; sudo ipmitool mc info 2>&1"''',
994: r'''sshpass -p "$DUT_PASS" ssh -o StrictHostKeyChecking=no $DUT_USER@$DUT_IP "sudo ipmitool raw 0x06 0x03 2>&1; sleep 30; sudo ipmitool mc info 2>&1"''',
998: r'''sshpass -p "$DUT_PASS" ssh -o StrictHostKeyChecking=no $DUT_USER@$DUT_IP "sudo ipmitool raw 0x06 0x07 2>&1"''',
999: r'''sshpass -p "$DUT_PASS" ssh -o StrictHostKeyChecking=no $DUT_USER@$DUT_IP "sudo ipmitool mc guid 2>&1"''',
1000: r'''sshpass -p "$DUT_PASS" ssh -o StrictHostKeyChecking=no $DUT_USER@$DUT_IP "sudo ipmitool raw 0x06 0x22 2>&1"''',
1001: r'''sshpass -p "$DUT_PASS" ssh -o StrictHostKeyChecking=no $DUT_USER@$DUT_IP "sudo ipmitool raw 0x06 0x24 2>&1"''',
1002: r'''sshpass -p "$DUT_PASS" ssh -o StrictHostKeyChecking=no $DUT_USER@$DUT_IP "sudo ipmitool raw 0x06 0x25 2>&1"''',
# --- Q-LIT: inner grep double quoted pattern -> single quotes (COM port/SOL) ---
833: r'''-- COM 0 -> COM port: set Linux console through COM port (ttyS0). Agent verifies the current console/kernel cmdline sshpass -p "$DUT_PASS" ssh -o StrictHostKeyChecking=no $DUT_USER@$DUT_IP "cat /proc/cmdline 2>&1; grep -E 'console' /etc/default/grub 2>&1"; setting/confirming the COM port (ttyS0) console output needs terminal access at the COM port (ttyS0) side.''',
834: r'''-- COM 0 -> SOL: set Linux console through SOL (ttyS0). Agent verifies the current console/kernel cmdline sshpass -p "$DUT_PASS" ssh -o StrictHostKeyChecking=no $DUT_USER@$DUT_IP "cat /proc/cmdline 2>&1; grep -E 'console' /etc/default/grub 2>&1"; setting/confirming the SOL (ttyS0) console output needs terminal access at the SOL (ttyS0) side.''',
# --- R5: OOB command wrongly wrapped in DUT-ssh + broken nested quotes -> move to agent-host ---
814: r'''-- Front Panel Power Button: agent powers the SUT on/off via OOB IPMI; the front-panel Power LED color (green/orange) is visual and the physical button press is operator-only (agent only drives the power command).
ipmitool -I lanplus -C 17 -H "$BMC_IP" -U "$BMC_USER" -P "$BMC_PASS" chassis power on 2>&1; sleep 5; ipmitool -I lanplus -C 17 -H "$BMC_IP" -U "$BMC_USER" -P "$BMC_PASS" chassis power status 2>&1''',
815: r'''-- Healthy LED: trigger/clear an injectable error via OOB IPMI and read SEL; the Healthy LED color (green=healthy/amber=fault) must be seen by a person.
ipmitool -I lanplus -C 17 -H "$BMC_IP" -U "$BMC_USER" -P "$BMC_PASS" sel elist 2>&1 | tail -20''',
987: r'''-- Get CPU current power consumption: read CPU/power sensors (OOB) + Redfish Chassis Power (agent-host curl).
ipmitool -I lanplus -C 17 -H "$BMC_IP" -U "$BMC_USER" -P "$BMC_PASS" sensor list 2>&1 | grep -iE "power"; echo ---redfish---; curl -s -k -u "$BMC_USER:$BMC_PASS" "https://$BMC_IP/redfish/v1/Chassis/1/Power" 2>&1 | jq -r ".PowerControl[]" 2>&1''',
988: r'''-- Get CPU capping/limit: read the current/max CPU power cap/limit (OOB sensor) + Redfish PowerLimit/PowerControl (agent-host curl).
ipmitool -I lanplus -C 17 -H "$BMC_IP" -U "$BMC_USER" -P "$BMC_PASS" sensor list 2>&1 | grep -iE "power"; echo ---redfish---; curl -s -k -u "$BMC_USER:$BMC_PASS" "https://$BMC_IP/redfish/v1/Chassis/1/Power" 2>&1 | jq -c ".PowerControl" 2>&1''',
990: r'''-- Get max CPU power capping/limit: read the current/max CPU power cap/limit (OOB sensor) + Redfish PowerLimit/PowerControl (agent-host curl).
ipmitool -I lanplus -C 17 -H "$BMC_IP" -U "$BMC_USER" -P "$BMC_PASS" sensor list 2>&1 | grep -iE "power"; echo ---redfish---; curl -s -k -u "$BMC_USER:$BMC_PASS" "https://$BMC_IP/redfish/v1/Chassis/1/Power" 2>&1 | jq -c ".PowerControl" 2>&1''',
}

wb = openpyxl.load_workbook(XLSX)
ws = wb[SHEET]
changed = 0
missing = []
for rownum, newcmd in FIX.items():
    code = ws.cell(row=rownum, column=1).value
    if code is None:
        missing.append(rownum)
        continue
    ws.cell(row=rownum, column=15, value=newcmd)
    changed += 1
if missing:
    print('MISSING rows:', missing)
print('changed', changed, 'of', len(FIX), 'on sheet', SHEET)
if changed != len(FIX):
    sys.exit('ERROR: not all rows updated')
wb.save(XLSX)
print('saved', XLSX)
