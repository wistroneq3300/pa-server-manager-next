# Closing verification: rebuild + cell diff vs closing_pre + mojibake + sheet counts.
import openpyxl, re, subprocess, sys

XLSX = 'data/REVISED_commands_merged_with_raw.xlsx'
BAK = 'data/REVISED_commands_merged_with_raw.xlsx.bak_rev02_closing_pre'
OUT = '/tmp/tests_closing.json'
SHEETS = ['Functionality', 'Reliability', 'Performance', 'Compatibility', 'Stability', '(No Main Function)']

# 1) rebuild
r = subprocess.run([sys.executable, 'scripts/build_testlib_json_xlsx.py', XLSX, OUT],
                   capture_output=True, text=True)
print(r.stdout[-400:])

# 2) cell diff vs closing_pre (expect 233, all col15)
import shutil, tempfile, os
_tmpbak = '/tmp/bak_closing_verify.xlsx'
shutil.copy(BAK, _tmpbak)
wbc = openpyxl.load_workbook(XLSX, data_only=True)
wbb = openpyxl.load_workbook(_tmpbak, data_only=True)
diff = []
for sn in SHEETS:
    sc = wbc[sn]
    sb = wbb[sn]
    for row in range(2, max(sc.max_row, sb.max_row) + 1):
        for col in range(1, max(sc.max_column, sb.max_column) + 1):
            if sc.cell(row=row, column=col).value != sb.cell(row=row, column=col).value:
                diff.append((sn, row, col))
from collections import Counter
print('diff vs closing_pre:', len(diff), '| by col:', dict(Counter(c for _, _, c in diff)),
      '| by sheet:', dict(Counter(s for s, _, _ in diff)))

# 3) mojibake scan (cyrillic / U+FFFD) across ai_commands + criteria, all sheets
bad = []
for sn in SHEETS:
    ws = wbc[sn]
    hdr = [c.value for c in ws[1]]
    ci = {h: i for i, h in enumerate(hdr)}
    for row in range(2, ws.max_row + 1):
        for colname in ('ai_commands', 'Criteria'):
            v = ws.cell(row=row, column=ci[colname] + 1).value
            if isinstance(v, str):
                if '\ufffd' in v or re.search(r'[\u0400-\u04ff]', v):
                    bad.append((sn, row, colname))
print('mojibake cells:', len(bad), bad[:5])

# 4) sdr fallback presence & <placeholder> residuals
nofb = []
ph = []
for sn in SHEETS:
    ws = wbc[sn]
    hdr = [c.value for c in ws[1]]
    ci = {h: i for i, h in enumerate(hdr)}
    for row in range(2, ws.max_row + 1):
        aic = ws.cell(row=row, column=ci['ai_commands'] + 1).value or ''
        if re.search(r'\bsensor\s+get\b', aic) and 'sdr elist' not in aic:
            nofb.append((sn, row))
        for mm in re.finditer(r'<([a-zA-Z_][a-zA-Z0-9_]*)>', aic):
            ph.append((sn, row, mm.group(1)))
print('sensor-get rows without sdr fallback:', len(nofb), nofb[:5])
print('<placeholder> residuals:', len(ph), ph[:5])
