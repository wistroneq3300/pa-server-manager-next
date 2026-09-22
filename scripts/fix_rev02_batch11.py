# Precise per-(sheet,row) ai_commands fix for Round-02 batch 11 (Functionality rows 2003-2202, item 2001-2200).
#
# Classes fixed (all ai_commands col15 unless noted):
#  Q-LIT      : inner unescaped double-quoted literal inside the outer ssh double-quoted string -> single-quote it.
#  Q-LIT+DBL : Q-LIT + collapse two sequential sshpass-ssh hops into one + single-quote the inner grep/echo literal.
#  R5         : OOB/BMC access wrongly wrapped in a DUT-ssh double hop (runs on the wrong host) -> agent-host direct.
#  WR         : command content does not match the Items/Procedure -> rewrite to the procedure's real logic.
#  VERDICT    : col13 NO/PARTIAL -> YES when the ai_commands was a copy-error text and the real test is the
#               same read-only sensor read as the sibling rows in the same block (mirrors batch10 row 1843).
#  PKG        : col14 ai_packages_needed corrected to the real read-only tool when col13 flipped.
#
# Pure ASCII in the cell text where the original is ASCII. Row-located (not code-located). python-openpyxl.
#
# Construction convention for a cell string:
#   note_lines joined by chr(10). The SSH remote is written as  ss(<remote>)  =  SSH + DQ + <remote> + DQ.
#   <remote> carries only single-quoted patterns plus, where a backslash must SURVIVE into the cell
#   (e.g. \$OAM for a remote loop var, or \" for JSON inside the remote), the literal backslash is
#   emitted via the BS (chr(92)) constant so the source stays legible and unambiguous.
import openpyxl, sys

XLSX = 'data/REVISED_commands_merged_with_raw.xlsx'
SHEET = 'Functionality'
NL = chr(10)
BS = chr(92)
DQ = chr(34)

SSH = 'sshpass -p "$DUT_PASS" ssh -o StrictHostKeyChecking=no $DUT_USER@$DUT_IP'
OOB = 'ipmitool -I lanplus -C 17 -H "$BMC_IP" -U "$BMC_USER" -P "$BMC_PASS"'
BMC_SSH = 'sshpass -p "$BMC_PASS" ssh -o StrictHostKeyChecking=no "$BMC_USER"@"$BMC_IP"'


def ss(remote):
    return SSH + DQ + remote + DQ


def bmc_ss(remote):
    return BMC_SSH + DQ + remote + DQ


FIX = {}
CAN = {}
PKG = {}

# ---------- Q-LIT : inner double-quote in ssh remote -> single-quote ----------
FIX[2049] = ('-- XGMI 4 Point Screen Test: XGMI 4-point screen test via amdxio. Agent runs the link status then the 4pt screen test on all links'
             + NL + ss("sudo ./amdxio -xgmi -linkstatus 2>&1 | tail -n 20; echo ---4pt---; sudo ./amdxio -margin -xgmi -all -ber=10 -errcnt=5 -4pt -phaseleftrange=2:0.2 -phaserightrange=2:0.2 -dactoprange=4:0.4 -dacbotrange=4:0.4 2>&1 | tail -n 30 2>&1")
             + NL + '-- amdxio is a vendor tool (operator-provided on the DUT).')
FIX[2066] = ('-- Post-Boot Stress: after each reboot run an OS/GPU health sweep on the AMD SVM. Agent can run the per-boot checks (dmesg errors, amdgpu state)'
             + NL + ss("dmesg -l err 2>&1 | tail -n 40; echo ---; rocm-smi 2>&1 | head -n 20 2>&1")
             + NL + '-- automating the reboot loop itself needs a reboot controller (operator).')
FIX[2105] = ('-- SW.4_DOCA Performance: run the NVNetPerf test suites to qualify DOCA networking performance. Agent can pre-check the DOCA/network stack'
             + NL + ss("docactl --version 2>&1; ethtool -i $DUT_NIC 2>&1 | head -n 6; ibstat 2>&1 | grep -iE 'state|link' 2>&1")
             + NL + '-- the full NVNetPerf suite is a long benchmark needing the vendor test setup (operator provides DOCA SDK + NVNetPerf).')
FIX[2123] = ('-- NT.2_Network NVQual: run the NVIDIA NVQual network test on the BlueField.'
             + NL + ss("nvqual -h 2>&1 | head -n 15 2>&1")
             + NL + '-- NVQual is a vendor tool (operator install + license/test topology); agent provides the launch wrapper + collects output, then runs the specific NVQual network test per the NVIDIA docs.')
FIX[2125] = ('-- NT.3_Network NVQual: run the NVIDIA NVQual network test on the BlueField.'
             + NL + ss("nvqual -h 2>&1 | head -n 15 2>&1")
             + NL + '-- NVQual is a vendor tool (operator install + license/test topology); agent provides the launch wrapper + collects output, then runs the specific NVQual network/test per the NVIDIA docs.')
FIX[2149] = ('-- US.1_USB2.0 Interface between HMC and BMC: verify the HMC-BMC USB2.0 link enumerates to high speed. Agent checks the USB tree for the HMC management port'
             + NL + ss("lsusb -t 2>&1; dmesg 2>&1 | grep -iE 'usb 2-?[0-9]|SuperSpeed|high-speed usb' | tail -n 20 2>&1"))
FIX[2151] = ('-- SY.2_CPU Motherboard: validate the CPU motherboard per CPU-vendor guidance; agent runs the generic platform health sweep'
             + NL + ss("lscpu 2>&1 | grep -E 'Model name|Socket|Core'; dmidecode -t 4 2>&1 | grep -iE 'Version|Manufacturer'; lspci 2>&1 | grep -iE 'VGA|Processing accelerators' 2>&1"))
FIX[2154] = ('-- SW.1_System software: run the NVSSVT system-software qualification suites (L0/L1). The agent can collect the software/OS inventory that NVSSVT validates'
             + NL + ss("nvidia-smi 2>&1 | head -n 12; echo ---; uname -r; cat /etc/os-release 2>&1 | grep -E 'PRETTY_NAME' 2>&1")
             + NL + '-- the full NVSSVT suite is a multi-test vendor-run needing the NVSSVT package (operator provides).')
FIX[2155] = ('-- SW.2_Server RAS Capability: run the NVRAS server-RAS-capability suites. Agent first collects the RAS-relevant OS state'
             + NL + ss("rasdaemon --help 2>&1 | head -n 5; ls /sys/devices/system/edac/ 2>&1; dmesg 2>&1 | grep -iE 'ras|edac' | tail -n 10 2>&1")
             + NL + '-- the full NVRAS validation needs the NVIDIA NVRAS tool suites (operator provides).')
FIX[2166] = ('-- UA.1_UART CPU to USB-C: verify UART CPU-to-USB-C connectivity. Agent can enumerate the USB-C serial adapters'
             + NL + ss("dmesg 2>&1 | grep -iE 'cp210x|ftdi|usb 1-' | tail -n 15; ls /dev/ttyUSB* /dev/ttyACM* 2>&1")
             + NL + '-- a functional send/receive traffic loop needs a physical USB-C connection (operator).')
FIX[2167] = ('-- UA.2_UART BMC to USB-C: verify UART BMC-to-USB-C connectivity. Agent lists the serial devices under BMC'
             + NL + bmc_ss("ls /dev/ttyS* /dev/ttyUSB* 2>&1; dmesg 2>&1 | grep -iE 'serial|ttyS|usb' | tail -n 12 2>&1")
             + NL + '-- a send/receive loop needs an operator on the USB-C console.')
FIX[2168] = ('-- UA.3_UART HMC to USB-C: verify UART HMC-to-USB-C. Agent checks the HMC serial/console device state'
             + NL + ss("ls /dev/ttyACM* /dev/ttyUSB* 2>&1; dmesg 2>&1 | grep -iE 'acm|usb' | tail -n 12 2>&1")
             + NL + '-- a bidirectional UART traffic test needs a physical terminal on the USB-C connector (operator).')


def USB(rn, prefix):
    return ('-- %s: verifying the USB hub/port link + downstream devices is a physical check (the hub/port is on the BMC/HPM/PCI board); agent reads the OS USB topology' % prefix
            + NL + ss("lsusb 2>&1 | grep -iE 'hub|usb' | head -n 40 2>&1")
            + NL + '-- once the BMC/port link is active; the physical connection/link-up is operator-observed.')


FIX[2169] = USB(2169, 'US.1_USB BMC to HMC USB Hub functional')
FIX[2170] = USB(2170, 'US.2_USB BMC to HPM USB Hub functional')
FIX[2171] = USB(2171, 'US.3_USB BMC to Bay B CX9 USB Hub functional')
FIX[2172] = USB(2172, 'US.4_USB BMC to BF4 USB Hub functional')
FIX[2173] = USB(2173, 'US.5_USB CPU to USB port on Front IO board functional')
FIX[2174] = USB(2174, 'US.6_USB CPU to USB controller (ETH SW)  functional')

FIX[2194] = ('-- Vendor ID: read CPU vendor/model/architecture + dmidecode type-4 manufacturer/version/family.'
             + NL + ss("lscpu 2>&1 | grep -iE 'vendor|model name|architecture'; dmidecode -t 4 2>/dev/null | grep -iE 'manufacturer|version|family' | head 2>&1"))


def MODS(rn, items):
    return ('-- %s: run the test via MODS (vendor platform tool, operator installs + provides the run recipe).' % items
            + NL + ss("mods -h 2>&1 | head -n 20 2>&1")
            + NL + '-- then run the matching MODS test module per the vendor MODS user guide (operator-provided recipe).')


FIX[2198] = MODS(2198, 'CPUDVFS Test')
FIX[2199] = MODS(2199, 'CPU Thermal Stress Test')
FIX[2200] = MODS(2200, 'USB2.0 MCTP Function')
FIX[2201] = MODS(2201, 'NVMe/E1.S Read/Write test')

# ---------- Q-LIT : BMC-remove-management rows 2050-2055 (DUT-side BMC VNIC, stay on DUT) ----------
# 192.168.31.1 = DUT-side BMC VNIC -> reachable only from the DUT host, so NOT R5. Inner
# -u "$.." double-quote -> single-quote; remote loop vars that must expand on the DUT -> \$OAM.
BMV = "'$BMC_USER:$BMC_PASS'"
FIX[2050] = ('-- Telemetry Data for UBB/OAM: query UBB/OAM Redfish TelemetryService metric report over the BMC VNIC (192.168.31.1, reachable from the DUT host), then check the readings are within the lower/upper limit (non-NULL, non-zero).'
             + NL + ss('curl -s -k -u ' + BMV + ' http://192.168.31.1/redfish/v1/TelemetryService/MetricReports/All 2>&1'))
FIX[2051] = ('-- Collect UBB FW Info and Update UBB FW: collect UBB FW info via Redfish FirmwareInventory (BMC VNIC), then update UBB FW; the flash/update step needs the new FW image + upload action (operator).'
             + NL + ss('curl -s -k -u ' + BMV + ' http://192.168.31.1/redfish/v1/UpdateService/FirmwareInventory 2>&1; echo ---retimer_active---; curl -s -k -u ' + BMV + ' http://192.168.31.1/redfish/v1/UpdateService/FirmwareInventory/retimer_active 2>&1'))
FIX[2052] = ('-- Health Check on UBB and OAM: health-check UBB and OAM0-7 via Redfish Chassis (BMC VNIC). Load amdgpu first.'
             + NL + ss('modprobe amdgpu 2>&1; curl -s -k -u ' + BMV + ' http://192.168.31.1/redfish/v1/Chassis/UBB 2>&1; for OAM in 0 1 2 3 4 5 6 7; do curl -s -k -u ' + BMV + ' http://192.168.31.1/redfish/v1/Chassis/OAM_' + BS + '$OAM; echo; done 2>&1'))
FIX[2053] = ('-- SMC Reset and Set DateTime: force-reset the SMC and set DateTime via Redfish Manager/AMC (BMC VNIC); setting DateTime is a state write (operator approves).'
             + NL + ss('curl -s -k -u ' + BMV + " -X POST -d '{" + BS + DQ + 'ResetType' + BS + DQ + ':' + BS + DQ + 'ForceRestart' + BS + DQ + "}' http://192.168.31.1/redfish/v1/Managers/AMC/Actions/Manager.Reset 2>&1; ping -c 4 192.168.31.1 2>&1; curl -s -k -u " + BMV + ' http://192.168.31.1/redfish/v1/Managers/AMC 2>&1'))
FIX[2054] = ('-- Log Collection: collect SMC/system logs via Redfish LogServices (BMC VNIC): list LogServices then dump EventLog entries.'
             + NL + ss('curl -s -k -u ' + BMV + ' http://192.168.31.1/redfish/v1/Systems/UBB/LogServices 2>&1; curl -s -k -u ' + BMV + ' http://192.168.31.1/redfish/v1/Systems/UBB/LogServices/EventLog/Entries >> event.log 2>&1; wc -l event.log 2>&1'))
FIX[2055] = ('-- Power Management: manage the Power Limit Watts for UBB/OAM via OOB Redfish (BMC VNIC) plus in-band ROCm; first read the current EnvironmentMetrics SetPoint for each OAM; the actual power-limit set uses ROCm/power-cap tooling and is a state write (operator).'
             + NL + ss('for OAM in 0 1 2 3 4 5 6 7; do curl -s -k -u ' + BMV + ' http://192.168.31.1/redfish/v1/Chassis/OAM_' + BS + '$OAM/EnvironmentMetrics; echo; done 2>&1'))

# ---------- R5 : BMC Dimm-usage row 2092 reached the BMC shell via a DUT-host double hop ----------
# Move to agent-host -> BMC shell direct.
FIX[2092] = ('-- Check BMC Dimm usage rate under idle: log into the BMC shell and sample process memory to confirm no abnormal growth (agent-host -> BMC shell, OOB).'
             + NL + bmc_ss("top -b -n 1 2>&1 | grep -Ei 'redfish' | head -n 5; free -m 2>&1; cat /proc/meminfo 2>&1 | head -n 5 2>&1"))

# ---------- Q-LIT + DBL-SSH : BIOS C/P-state rows -> collapse two hops into one + single quote ----------
FIX[2043] = ('-- C-state (Core States): verify CPU C-states are enabled/working. C-state config is a BIOS CPU-setting (human); the OS residency read is automated.'
             + NL + ss("lscpu 2>&1 | grep -i -E 'c-state|model name'; grep . /sys/devices/system/cpu/cpu0/cpuidle/state*/name 2>&1; turbostat -q -n 1 2>&1 | tail -n 5 2>&1"))
FIX[2044] = ('-- P-state (Performance States): verify CPU P-states/OS performance states. P-state config is a BIOS CPU-setting; the OS governor read is automated.'
             + NL + ss("lscpu | grep -i -E 'MHz|min|max'; cpupower frequency-info 2>&1 | grep -iE 'driver|governor|hardware limits' | head 2>&1"))

# ---------- WR : command content / sensor name does not match Items/Procedure ----------
FIX[2012] = ('-- Sensor Check - PCH: read the PCH sensor via ipmitool (OOB, from the agent host) and confirm name/description/reading match SPEC.'
             + NL + OOB + " sensor get 'Sensor Check - PCH' 2>&1")
FIX[2028] = OOB + " sensor get 'Power Reading Check - PWR' 2>&1"
FIX[2039] = ('-- Check Status - GB: read the GB status sensor via ipmitool (OOB, from the agent host) and confirm name/description/reading match SPEC.'
             + NL + OOB + " sensor get 'Check Status - GB' 2>&1")
FIX[2045] = ('-- Flash PCIe Switch Firmware: run the vendor PCIe-switch FW flash tool (g4Xdiagnostics) with the FW image, then verify the updated FW version via the tool/lspci.'
             + NL + ss("unzip -o /opt/g4Xd/g4Xdiagnostics_linux.zip -d /opt/g4Xd/ >/dev/null 2>&1; cd /opt/g4Xd && sudo ./g4Xdiag flash " + '$' + '{SW_FW_IMAGE:?operator provides the PCIe-Switch FW image}' + ' 2>&1 | tail -n 40 2>&1')
             + NL + '-- g4Xdiagnostics is a vendor tool (operator provides the tool + FW image on the DUT).')
FIX[2202] = ('-- Telemetry INA(INA3221) / OVRM(NCP45495XMNTWG): read the telemetry current/power sensors via ipmitool (OOB, from the agent host) and confirm the INA/OVRM readings match SPEC.'
             + NL + OOB + " sdr list 2>&1 | grep -Ei 'INA|OVRM|PWR' | head -n 60")

# ---------- VERDICT + PKG flips (read-only sensor read, mirror batch10 1843) ----------
CAN[2012] = 'YES'; PKG[2012] = 'ipmitool'
CAN[2039] = 'YES'; PKG[2039] = 'ipmitool'


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else XLSX
    dry_only = (sys.argv[2] == 'dry') if len(sys.argv) > 2 else False
    wb = openpyxl.load_workbook(xlsx)
    if SHEET not in wb.sheetnames:
        print('ERROR: sheet %r not in workbook' % SHEET); sys.exit(2)
    ws = wb[SHEET]
    n_cmd = n_can = n_pkg = 0
    for r, v in sorted(FIX.items()):
        cur = ws.cell(row=r, column=15).value
        if cur is not None and str(cur) == v:
            print('row %d: unchanged (already correct)' % r); continue
        if dry_only:
            print('--- row %d new ai_commands ---' % r); print(v); continue
        ws.cell(row=r, column=15).value = v
        n_cmd += 1
    if dry_only:
        wb.close(); print('DRY run -- NOT saving.'); return
    for r, v in sorted(CAN.items()):
        ws.cell(row=r, column=13).value = v
        n_can += 1
    for r, v in sorted(PKG.items()):
        ws.cell(row=r, column=14).value = v
        n_pkg += 1
    wb.save(xlsx)
    wb.close()
    print('applied  col15 ai_commands: %d   col13 ai_can_execute: %d   col14 ai_packages_needed: %d' % (n_cmd, n_can, n_pkg))

    wb2 = openpyxl.load_workbook(xlsx, read_only=True)
    ws2 = wb2[SHEET]
    issues = []
    for r in sorted(FIX.keys()):
        v = ws2.cell(row=r, column=15).value
        s = '' if v is None else str(v)
        if 'ssh ' in s or 'sshpass' in s:
            if s.count('"') % 2 != 0:
                issues.append((r, 'odd double-quote count %d' % s.count('"')))
            if r in (2043, 2044) and s.count('sshpass') >= 2:
                issues.append((r, 'still has %d nested sshpass' % s.count('sshpass')))
    wb2.close()
    if issues:
        print('post-audit ISSUES:'); [print('   row', r, why) for r, why in issues]
    else:
        print('post-audit issues: NONE')


if __name__ == '__main__':
    main()
