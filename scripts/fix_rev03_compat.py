#!/usr/bin/env python3
"""Round-3 Batch-3 fixes (2026-09-18) - Compatibility sheet (only r224).

r224 (Wistron-AMD GPU-00217-V002): bare `$FRUfile` -> ${FRU_FILE:?} vendor slot
(R7). col15 + col16 both reference the var. Everything else already clean
(round-1 + closing/run2/run3 passes).
"""
import sys
import openpyxl

XLSX = 'data/REVISED_commands_merged_with_raw.xlsx'
DRY = '--dry' in sys.argv

wb = openpyxl.load_workbook(XLSX)
C = wb['Compatibility']


def setcell(row, col, old, new):
    cur = C.cell(row=row, column=col).value
    if cur is None:
        cur = ''
    if cur != old:
        raise AssertionError(f"Compatibility r{row} col{col} mismatch:\n  got: {cur!r}\n  exp: {old!r}")
    if not DRY:
        C.cell(row=row, column=col).value = new
    print(f"Compatibility r{row:3d} col{col:2d} {'DRY' if DRY else 'SET'} ({len(old)} -> {len(new)} chars)")


setcell(224, 15,
 '-- AGT FRU flash needs the user-specified FRU data file on DUT; run `agt fru flash -f $FRUfile` then read back to confirm the flashed FRU matches the file.',
 '-- AGT FRU flash needs the user-specified FRU data file on DUT; run `agt fru flash -f ${FRU_FILE:?operator provides the FRU data file path on the DUT}` then read back to confirm the flashed FRU matches the file.')
setcell(224, 16,
 'return AGT FRU flash result + readback; user confirms FRU data == $FRUfile.',
 'return AGT FRU flash result + readback; user confirms FRU data == ${FRU_FILE:?operator provides the FRU data file path on the DUT}.')

if not DRY:
    wb.save(XLSX)
    print('saved:', XLSX)
else:
    print('DRY-RUN - not saved')
