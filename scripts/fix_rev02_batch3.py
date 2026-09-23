# Precise per-(sheet,row) ai_commands fix for Round-02 batch 3 (Functionality rows 402-601).
# Q-LIT: inner grep double-quoted pattern -> single quotes inside ssh.
# E-DQ: inner echo double-quoted string -> single quotes inside ssh.
# Pure ASCII only in this script. Uses python-openpyxl.
import openpyxl, sys

XLSX='data/REVISED_commands_merged_with_raw.xlsx'
SHEET='Functionality'
FIX={
474: r'''# Check NUMA node configuration after operator sets BIOS NPS (NPS1/2/4)
sshpass -p "$DUT_PASS" ssh -o StrictHostKeyChecking=no $DUT_USER@$DUT_IP "lscpu 2>&1 | grep -E '^NUMA|Socket|Core|On-line'; echo '---node memory---'; numactl --hardware 2>&1 | grep -i node"''',
476: r'''# Devices Firmware Versions SSD: compare OS-side SSD FW (BIOS-side read is operator)
sshpass -p "$DUT_PASS" ssh -o StrictHostKeyChecking=no $DUT_USER@$DUT_IP "nvme id-ctrl /dev/nvme0n1 2>&1 | grep -Ei 'fr |mn |sn '; hdparm -I /dev/sda 2>&1 | grep -Ei 'Firmware Revision'"''',
480: r'''# SMBIOS version check (2.2): read SMBIOS version + confirm stable after reboot
sshpass -p "$DUT_PASS" ssh -o StrictHostKeyChecking=no $DUT_USER@$DUT_IP "sudo dmidecode -t 0 2>&1 | grep -Ei 'SMBIOS|BIOS Revision|Release Date'; sudo dmidecode --version"''',
}

wb=openpyxl.load_workbook(XLSX)
ws=wb[SHEET]
changed=0; missing=[]
for rownum,newcmd in FIX.items():
    code=ws.cell(row=rownum, column=1).value
    if code is None:
        missing.append(rownum); continue
    ws.cell(row=rownum, column=15, value=newcmd)
    changed+=1
if missing: print('MISSING rows:',missing)
print('changed',changed,'of',len(FIX),'on sheet',SHEET)
if changed!=len(FIX):
    sys.exit('ERROR: not all rows updated')
wb.save(XLSX)
print('saved',XLSX)
