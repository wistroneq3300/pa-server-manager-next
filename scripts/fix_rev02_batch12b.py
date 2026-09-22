# batch12 FOLLOW-UP sync: rows whose ai_commands (col15) was rewritten in batch 12 still carried
# the OLD text in ai_logs_output (col16) and risk (col17). This fixes ONLY those two columns so
# the WR/FAKE rewrites are internally consistent (col15 <-> col16/col17). Col13/14/15 untouched.
#
# Guard: each row MUST currently contain the batch-12 new ai_commands start marker, else ABORT.
import openpyxl, sys

XLSX = 'data/REVISED_commands_merged_with_raw.xlsx'
SHEET = 'Functionality'

# keys: row -> (expected_col15_marker, new_col16, new_col17)
# new_col17 None means "leave as-is", '' means "clear to None"
FIX = {
    2203: (
        '-- USB2/3 Function Check: verify the USB hub/device (MCU) enumerates',
        'return the full lsusb -t / lsusb output (hub/device enumeration + negotiated link speed 480M or 5Gbps) so the user confirms the USB hub/device (MCU) enumerates at the expected speed.',
        '',
    ),
    2209: (
        '-- AC Cycle - Speed&Link: the operator performs the AC power cycle',
        "return the lspci link speed/width recorded before and after each AC cycle so the user confirms all PCIe speeds/link statuses are normal after every cycle.",
        'RISK: AC power cycling interrupts the SUT; the physical cycle + overnight run is operator-setup (agent only records lspci pre/post).',
    ),
    2241: (
        '-- Add HTTPBasicAuth - Redfish: verify the HTTPBasicAuth toggle',
        'return the baseline GET response (expect 200), the PATCH response for HTTPBasicAuth=Disabled, the re-GET (expect 401), and the post factory-reset GET so the user confirms the auth toggle and restore.',
        'RISK: disabling basic auth can lock out management access; the factory reset restores it - operator approve before the state change.',
    ),
    2272: (
        '-- TEMP_HIB_PEX: read the TEMP_HIB_PEX sensor via ipmitool',
        "return the full OOB `ipmitool sensor get 'TEMP_HIB_PEX'` output (name / description / reading) so the user confirms the sensor reads at the expected value.",
        None,
    ),
    2273: (
        '-- MLPerf - Retinanet Offline: pre-check the SUT inventory',
        'return the inventory pre-check + the MLPerf retinanet/offline benchmark result (latency/throughput) so the user confirms the score meets the MLCommons v5.0 reference.',
        None,
    ),
    2277: (
        '-- Release note: verify the FW package release note',
        'return the release-note review checklist (general information, FW version, new features/defects fixed, known issues) + the confirmed FW version; human document review, no code output.',
        None,
    ),
    2287: (
        '-- TH.1_Thermal Stress Test: run a full-power thermal stress',
        "return the thermal stress run log + the OOB thermal sensor reads (ipmitool sensor get temp) so the user confirms no thermal warnings/faults were logged.",
        'RISK: full-power thermal stress runs the DUT hot; the worst-case ambient / thermal chamber is operator setup.',
    ),
    2290: (
        '-- I3C - CPU1: verify the I3C bus to CPU1',
        'return the BMC-console I3C probe/device enumeration output so the operator confirms the I3C bus to CPU1 is connected/functional (vendor i3c-tools + exact probe steps are operator-provided).',
        None,
    ),
    2291: (
        '-- I3C - CPU1: verify the I3C bus to CPU1',
        'return the BMC-console I3C probe/device enumeration output so the operator confirms the I3C bus to CPU1 is connected/functional (vendor i3c-tools + exact probe steps are operator-provided).',
        None,
    ),
}

wb = openpyxl.load_workbook(XLSX)
ws = wb[SHEET]
hdr = [c.value for c in ws[1]]
ci = {h: i for i, h in enumerate(hdr)}
assert 'ai_commands' in hdr and 'ai_logs_output' in hdr and 'risk' in hdr
CC15 = ci['ai_commands'] + 1
CC16 = ci['ai_logs_output'] + 1
CC17 = ci['risk'] + 1

DRY = len(sys.argv) > 1 and sys.argv[1] == 'dry'
changed = 0
for row, (marker, new16, new17) in sorted(FIX.items()):
    c15 = ws.cell(row=row, column=CC15).value or ''
    if marker not in c15:
        sys.exit('ABORT row %d: col15 does not contain %r' % (row, marker))
    changed += 1
    ws.cell(row=row, column=CC16).value = new16
    if new17 is not None:
        ws.cell(row=row, column=CC17).value = (None if new17 == '' else new17)
    print('row %d: col16+col17 synced' % row)

print('total rows synced (incl. col16-only or col16+col17):', changed)
if DRY:
    print('DRY - not saved')
else:
    wb.save(XLSX)
    print('saved', XLSX)
