#!/usr/bin/env python3
"""verify_review.py - zero-regression gate for the test-library review workflow.

Usage:
    python3 verify_review.py <xlsx> <backup_xlsx> [--build "<build cmd>"] [--out <json>]

Reports: cell diff vs backup (per-sheet / per-col), shared-command recount,
mojibake (cyrillic / U+FFFD), literal '%s' remnants, sensor-get without sdr
fallback, nested/double sshpass + inner-quote (Q-LIT) candidates, and raw
<placeholder> tags in ai_commands. Pure-ASCII; needs only openpyxl.
"""
import sys, re, shutil, argparse, subprocess
from collections import Counter, defaultdict
import openpyxl

SHEETS = ['Functionality', 'Reliability', 'Performance',
          'Compatibility', 'Stability', '(No Main Function)']
MOJI = re.compile(r'[\u0400-\u04ff]')


def load(path):
    return openpyxl.load_workbook(path, data_only=True)


def colidx(ws, name):
    hdr = [c.value for c in ws[1]]
    try:
        return hdr.index(name) + 1
    except ValueError:
        return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsx')
    ap.add_argument('backup')
    ap.add_argument('--build', default=None, help='optional build cmd string')
    ap.add_argument('--out', default=None)
    a = ap.parse_args()

    if a.build:
        r = subprocess.run(a.build, shell=True, capture_output=True, text=True)
        print('BUILD rc=%d' % r.returncode)
        print(r.stdout[-400:])
        print(r.stderr[-200:])

    bak_tmp = '/tmp/_verify_review_backup.xlsx'
    shutil.copy(a.backup, bak_tmp)
    wc = load(a.xlsx)
    wb = load(bak_tmp)

    diff = []
    for sn in SHEETS:
        sc, sb = wc[sn], wb[sn]
        for row in range(2, max(sc.max_row, sb.max_row) + 1):
            for col in range(1, max(sc.max_column, sb.max_column) + 1):
                if sc.cell(row=row, column=col).value != sb.cell(row=row, column=col).value:
                    diff.append((sn, row, col))
    print('DIFF vs backup: %d cells | by col: %s | by sheet: %s'
          % (len(diff), dict(Counter(c for _, _, c in diff)),
             dict(Counter(s for s, _, _ in diff))))

    sh = defaultdict(list)
    for sn in SHEETS:
        if sn == 'Functionality':
            continue
        ws = wc[sn]
        ci = colidx(ws, 'ai_commands')
        for row in range(2, ws.max_row + 1):
            v = ws.cell(row=row, column=ci).value or ''
            if v.strip():
                sh[v].append((sn, row))
    shared_rows = sum(len(v) for v in sh.values() if len(v) >= 2)
    shared_sets = sum(1 for v in sh.values() if len(v) >= 2)
    print('SHARED-COMMAND rows: %d sets: %d (Functionality excluded)'
          % (shared_rows, shared_sets))

    bad = []
    for sn in SHEETS:
        ws = wc[sn]
        ci_a = colidx(ws, 'ai_commands')
        ci_c = colidx(ws, 'Criteria')
        for row in range(2, ws.max_row + 1):
            for ci, name in ((ci_a, 'ai_commands'), (ci_c, 'Criteria')):
                if ci is None:
                    continue
                v = ws.cell(row=row, column=ci).value
                if isinstance(v, str) and ('\ufffd' in v or MOJI.search(v)):
                    bad.append((sn, row, name))
    print('MOJIBAKE (cyrillic/U+FFFD): %d %s' % (len(bad), bad[:3]))

    pct = []
    for sn in SHEETS:
        ws = wc[sn]
        ci = colidx(ws, 'ai_commands')
        for row in range(2, ws.max_row + 1):
            v = ws.cell(row=row, column=ci).value or ''
            if '%s %s' in v or '-- %s:' in v:
                pct.append((sn, row))
    print('LITERAL %%s remnants: %d %s' % (len(pct), pct[:5]))

    nofb = []
    for sn in SHEETS:
        ws = wc[sn]
        ci = colidx(ws, 'ai_commands')
        for row in range(2, ws.max_row + 1):
            v = ws.cell(row=row, column=ci).value or ''
            if re.search(r'\bsensor\s+get\b', v) and 'sdr elist' not in v:
                nofb.append((sn, row))
    print('sensor-get WITHOUT sdr fallback: %d %s' % (len(nofb), nofb[:5]))

    dbl = []
    for sn in SHEETS:
        ws = wc[sn]
        ci = colidx(ws, 'ai_commands')
        for row in range(2, ws.max_row + 1):
            v = ws.cell(row=row, column=ci).value or ''
            if v.count('sshpass') > 1:
                dbl.append(('DBLSSH', sn, row))
            m = re.search(r'ssh[^\n"]*"([^"]*)"', v)
            if m and '"' in m.group(1):
                dbl.append(('QLIT', sn, row))
    print('nested-ssh / inner-quote candidates: %d %s' % (len(dbl), dbl[:15]))

    ph = []
    for sn in SHEETS:
        ws = wc[sn]
        ci = colidx(ws, 'ai_commands')
        for row in range(2, ws.max_row + 1):
            v = ws.cell(row=row, column=ci).value or ''
            for mm in re.finditer(r'<([a-zA-Z_][a-zA-Z0-9_]*)>', v):
                ph.append((sn, row, mm.group(1)))
    print('<placeholder> tags in ai_commands: %d %s' % (len(ph), ph[:5]))

    if a.build:
        ok = not bad and not pct and not nofb
        print('GATE: moji=%d pct=%d nofb=%d -> %s'
              % (len(bad), len(pct), len(nofb), 'PASS' if ok else 'REVIEW'))
    else:
        print('GATE: (no build run) diff=%d moji=%d pct=%d -> describe diff per-sheet/per-col in handoff'
              % (len(diff), len(bad), len(pct)))


if __name__ == '__main__':
    main()
